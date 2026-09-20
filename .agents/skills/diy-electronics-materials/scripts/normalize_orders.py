#!/usr/bin/env python3
"""Normalize sanitized LCSC captures and Taobao official exports for materials.py."""
import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urlsplit, urlunsplit


TAOBAO_TERMS = tuple(x.casefold() for x in """
元器件 芯片 集成电路 IC MCU 单片机 开发板 核心板 模块 模组 套件
ESP ESP32 STM32 Arduino M5Stack 树莓派 Raspberry Pico RISC-V FPGA
电阻 排阻 电容 电感 磁珠 二极管 三极管 MOS MOSFET 晶振 光耦 继电器
整流桥 稳压 运放 放大器 驱动器 保险丝 压敏 热敏 电位器 编码器 可控硅
TVS LDO DC-DC ADC DAC EEPROM Flash SD-NAND SMD DIP BGA QFN SOP 贴片 直插
传感器 光敏 霍尔 红外 毫米波 马达 电机 舵机 振动器
开关 按钮 按键 连接器 接插件 端子 排针 排母 插座 插头 排线 电子线
杜邦线 线束 FPC FFC PCB 面包板 洞洞板 电路板 硅胶线 同轴线 转接线
电池 锂电 电源 适配器 散热片 天线 GPS GNSS LTE 4G 蓝牙 LoRa WiFi
摄像头 麦克风 硅咪 咪头 SPM1423 喇叭 扬声器 蜂鸣器 显示屏 液晶 OLED
USB Type-C HDMI UART TTL CAN RS485 JTAG SWD I2C SPI 探针 编程器 烧录器
下载器 示波器 万用表 逻辑分析仪 焊台 电烙铁 焊接 焊锡 助焊 吸锡
热风枪 镊子 斜口钳 剥线钳 压线钳 热缩管 维修垫 EchoEar
""".split())

FINISHED_GOODS = re.compile(
    r"手机壳|保护壳|耳机|电脑音箱|蓝牙音箱|电视机|充电宝|电子秤|电子烟|"
    r"电子琴|电子票|电子券|电饭|电冰箱|热水壶|暖手|吹风机|电风扇|"
    r"吊灯|吸顶灯|台灯|灯箱|照明灯|灯泡|筒灯|射灯|行车记录仪|"
    r"吸尘器|吸奶器|电子表|电子请柬|电子相册|兑换券|"
    r"软件|激活码|产品密钥|会员|账号|帐号|课程|教程|图书|书籍|"
    r"流量卡|日租卡|手机日租|老人机|智能手机|备用手机|Apple/苹果\s*iPhone|"
    r"iPhone 5s移动|iPhone5s 未激活|MP4|遥控飞机|电子蜡烛|电炸锅|电磁炉|"
    r"鼠标|键盘|电脑显示器|硬盘盒|相纸|打印纸|标签纸|头盔|玩具|"
    r"歌曲|音乐优盘|置物架|货架挂钩|展示架|面板安装螺钉|自行车|电动车|摩托车"
)
DEV_EXCEPTIONS = re.compile(r"开发|模块|模组|套件|核心板|ESP|STM|M5Stack|芯片|单片机|PCB", re.I)


def clean(value):
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ")).strip()


def clean_url(value):
    if not value:
        return ""
    u = urlsplit(value)
    query = parse_qs(u.query)
    kept = {k: v for k, v in query.items() if k in {"id"}}
    return urlunsplit((u.scheme or "https", u.netloc, u.path, urlencode(kept, doseq=True), ""))


def quantity(value):
    m = re.search(r"(\d+(?:\.\d+)?)\s*([^\d\s]+)?", clean(value))
    if not m:
        return None, "未知"
    n = float(m.group(1))
    if n.is_integer():
        n = int(n)
    return n, (m.group(2) or "件")


def aliases(text):
    result = []
    for token in re.findall(r"(?i)\b[A-Z][A-Z0-9._/+\-]*\d[A-Z0-9._/+\-]*\b", text):
        token = token.strip("._/+-")
        if len(token) >= 3 and token.casefold() not in {x.casefold() for x in result}:
            result.append(token)
    return result[:20]


def category(text):
    rules = (
        (r"电阻|排阻", "电阻"), (r"电容", "电容"), (r"电感|磁珠", "电感与磁性器件"),
        (r"LED|二极管|整流桥|TVS", "二极管与LED"), (r"三极管|MOSFET|\bMOS\b", "晶体管"),
        (r"芯片|集成电路|单片机|MCU|LDO|DC-DC|ADC|DAC|运放|放大器|驱动器|EEPROM|Flash|SD-NAND", "集成电路"),
        (r"开发板|核心板|模块|模组|套件|ESP|STM32|Arduino|M5Stack|树莓派|Raspberry|FPGA", "模块与开发板"),
        (r"连接器|接插件|端子|排针|排母|插座|插头|排线|电子线|杜邦线|线束|FPC|FFC|Type-C|USB|HDMI|转接线|同轴线", "连接器与线材"),
        (r"电池|锂电", "电池"), (r"天线|GPS|GNSS|LTE|4G|LoRa|WiFi|蓝牙", "无线与天线"),
        (r"传感器|光敏|霍尔|红外|毫米波|激光|摄像头|麦克风|硅咪|咪头", "传感器"),
        (r"显示屏|液晶|OLED", "显示器件"), (r"马达|电机|舵机|振动器", "电机与执行器"),
        (r"开关|按钮|按键|编码器|继电器", "开关与继电器"),
        (r"电源|适配器|充电", "电源"), (r"焊|示波器|万用表|逻辑分析仪|镊子|钳|热风枪|维修垫|吸锡", "工具与耗材"),
        (r"PCB|面包板|洞洞板|电路板", "PCB与结构辅料"),
    )
    for pattern, name in rules:
        if re.search(pattern, text, re.I):
            return name
    return "电子配件"


def wanted_taobao(title, spec):
    text = f"{title} {spec}"
    folded = text.casefold()
    matched = False
    for term in TAOBAO_TERMS:
        if term == "4g":
            if re.search(r"(?<![A-Za-z0-9])4G(?![A-Za-z])", text):
                matched = True
                break
            continue
        if term.isascii():
            if re.search(rf"(?<![a-z]){re.escape(term)}(?![a-z])", folded, re.I):
                matched = True
                break
        elif term in folded:
            matched = True
            break
    if not matched:
        return False
    if FINISHED_GOODS.search(text) and not DEV_EXCEPTIONS.search(text):
        return False
    return True


def status(source, raw):
    raw = clean(raw)
    if source == "lcsc":
        return {"已发货": "shipped", "已完成": "completed", "已取消": "cancelled"}.get(raw, "unknown")
    if "交易成功" in raw or "充值成功" in raw:
        return "completed"
    if "已发货" in raw:
        return "shipped"
    if "关闭" in raw or "取消" in raw:
        return "cancelled"
    if "退款" in raw:
        return "refunded"
    if "待付款" in raw:
        return "unpaid"
    return "unknown"


def normalize_lcsc(raw_dir, observed):
    pages = []
    evidence_by_order = {}
    for page_no in range(1, 5):
        path = raw_dir / f"lcsc-page-{page_no}.json"
        for order in json.loads(path.read_text(encoding="utf-8")):
            pages.append(order)
            evidence_by_order[order["order_id"]] = str(path)
    details_path = raw_dir / "lcsc-order-details.json"
    details = {x["order_id"]: x for x in json.loads(details_path.read_text(encoding="utf-8"))}
    output = []
    for order in pages:
        products = details.get(order["order_id"], order)["products"]
        evidence = str(details_path if order["order_id"] in details else Path(evidence_by_order[order["order_id"]]))
        for idx, product in enumerate(products, 1):
            if "cells" in product:
                cells = product["cells"]
                info = clean(cells[3])
                parts = [clean(x) for x in re.split(r"\n+", cells[3]) if clean(x)]
                headline = parts[0] if parts else info
                model_pack = parts[1] if len(parts) > 1 else ""
                seller = parts[2] if len(parts) > 2 else ""
                model_parts = [clean(x) for x in model_pack.split("/", 1)]
                model = model_parts[0] if model_parts else ""
                package = model_parts[1] if len(model_parts) > 1 else ""
                q, unit = quantity(clean(cells[4]).split(" ", 1)[0])
                code = clean(product["code"])
                url = clean_url(product.get("url", ""))
            else:
                text = clean(product.get("text", ""))
                info = clean(text.split("编号：", 1)[0])
                headline = info
                model = clean(product.get("links", [""])[0] if product.get("links") else "")
                package = ""
                seller = ""
                m = re.search(r"编号：\s*(C\d+)", text, re.I)
                code = m.group(1).upper() if m else ""
                m = re.search(r"购买数量：\s*(\d+(?:\.\d+)?\s*[^商\s]*)", text)
                q, unit = quantity(m.group(1) if m else "")
                url = clean_url(product.get("url", ""))
            cat = clean(headline.split("/", 1)[0]) or "电子配件"
            spec = clean(headline.split("/", 1)[1] if "/" in headline else headline)
            a = aliases(f"{info} {model} {code}")
            output.append({
                "source": "lcsc", "account": "default", "order_id": clean(order["order_id"]),
                "line_id": f"{idx:03d}", "title": info, "quantity": q, "unit": unit,
                "status": status("lcsc", order.get("status", "")), "status_raw": clean(order.get("status", "")),
                "evidence": evidence, "observed_at": observed, "purchased_at": clean(order.get("date", "")),
                "seller": seller, "model": model, "lcsc_code": code, "spec": spec,
                "package": package, "category": cat, "product_url": url, "aliases": a,
            })
    return output


def normalize_taobao(raw_dir, observed):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("openpyxl is required to read Taobao official exports") from exc
    output = []
    counters = {}
    for path in sorted(raw_dir.glob("taobao-pages-*.xlsx")):
        ws = load_workbook(path, read_only=True, data_only=True).active
        meta = [None, None, None, None]
        for values in ws.iter_rows(min_row=2, values_only=True):
            values = list(values)
            for i in range(4):
                if values[i] is not None:
                    meta[i] = clean(values[i])
            order_id, purchased_at, raw_status, seller = meta
            title, url, spec, raw_q = (clean(x) for x in values[4:8])
            if not order_id or not title or not wanted_taobao(title, spec):
                continue
            normalized_status = status("taobao", raw_status)
            if normalized_status in {"cancelled", "unpaid", "refunded"}:
                continue
            counters[order_id] = counters.get(order_id, 0) + 1
            a = aliases(f"{title} {spec}")
            model = a[0] if a else ""
            q, unit = quantity(raw_q)
            output.append({
                "source": "taobao", "account": "default", "order_id": order_id,
                "line_id": f"{counters[order_id]:03d}", "title": title, "quantity": q,
                "unit": "件" if unit == "未知" else unit, "status": normalized_status,
                "status_raw": raw_status, "evidence": str(path), "observed_at": observed,
                "purchased_at": purchased_at, "seller": seller, "model": model,
                "spec": spec, "category": category(f"{title} {spec}"),
                "product_url": clean_url(url), "aliases": a, "notes": "",
            })
    return output


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("raw_dir", type=Path)
    ap.add_argument("output", type=Path)
    args = ap.parse_args()
    observed = datetime.now(timezone.utc).isoformat()
    rows = normalize_lcsc(args.raw_dir, observed) + normalize_taobao(args.raw_dir, observed)
    identities = {(r["source"], r["account"], r["order_id"], r["line_id"]) for r in rows}
    if len(identities) != len(rows):
        raise SystemExit("duplicate normalized order-line identity")
    args.output.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "output": str(args.output), "rows": len(rows),
        "lcsc_rows": sum(r["source"] == "lcsc" for r in rows),
        "taobao_rows": sum(r["source"] == "taobao" for r in rows),
        "orders": len({(r["source"], r["order_id"]) for r in rows}),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
